import flet as ft
from models.views.story import Story
from models.widget import Widget, WidgetView, WidgetSidebar
from flet_quill import FletQuill, FletQuillEditor, FletQuillToolbar
#from models.app import app
import math
import os
from utils.safe_string_checker import return_safe_name
import asyncio
import uuid
from styles.text_fields import TextField
from styles.snack_bar import SnackBar
from styles.context_menu_option import MenuOptionStyle
from dataclasses import dataclass, field

MANUSCRIPT_WIDTH = 820
MANUSCRIPT_HEIGHT = 1060
MANUSCRIPT_PADDING = 80
MANUSCRIPT_VERTICAL_MARGIN = 50
MANUSCRIPT_HORIZONTAL_MARGIN = 70


# Class that holds our text manuscript objects
@ft.observable
@dataclass
class Manuscript(Widget):

    tag: str = "manuscript" # Tag to differentiate this widget type
    
    manuscript_data: list = field(default_factory=list) # Holds the manuscript's text data as a list of json objects

    editor_dirty: bool = False  # True if the editor has unsaved changes

    def __post_init__(self):
        super().__post_init__()
        object.__setattr__(self, "quill_editor", None)   # Unannotated on purpose, so it isn't a dataclass field and never gets saved to json

    def attach_editor(self, editor, owner=None):
        object.__setattr__(self, "quill_editor", editor)
        object.__setattr__(self, "_editor_owner", owner)

    # Only the view that attached the editor may detach it, so a stale view's cleanup can't clear a newer view's editor
    def detach_editor(self, owner):
        if getattr(self, "_editor_owner", None) is owner:
            object.__setattr__(self, "quill_editor", None)
            object.__setattr__(self, "_editor_owner", None)
            
    # Checks if our manuscript is dirty, and saves it if it is
    async def save_file(self):
        if self.editor_dirty and self.quill_editor is not None:
            try:
                self.manuscript_data = await self.quill_editor.save()
                self.editor_dirty = False
                self.dirty = True
            except RuntimeError:
                pass    # Editor isn't on the page anymore, so it can't be read. Keep editor_dirty set
        await super().save_file()

    async def export(self, file_type: str="docx"):
        manuscript_data = self.manuscript_data
        paragraphs = self._parse_manuscript_ops(manuscript_data)

        if file_type == "docx":
            return self._export_docx(paragraphs)
        elif file_type == "pdf":
            return self._export_pdf(paragraphs)
        elif file_type == "txt":
            return self._export_txt(paragraphs)

        return None

    # Splits our quill delta ops (list of {'insert': str, 'attributes': dict}) into paragraphs of styled runs
    def _parse_manuscript_ops(self, ops: list) -> list:
        paragraphs = [[]]
        for op in ops:
            text = op.get('insert', "")
            attributes = op.get('attributes', {}) or {}
            if not isinstance(text, str):
                continue
            segments = text.split('\n')
            for idx, segment in enumerate(segments):
                if segment:
                    paragraphs[-1].append({'text': segment, **attributes})
                if idx < len(segments) - 1:
                    paragraphs.append([])
        if paragraphs and not paragraphs[-1]:
            paragraphs.pop()
        return paragraphs

    # Converts a '#RRGGBB'/'#AARRGGBB' hex color string into an (r, g, b) tuple
    def _hex_to_rgb(self, color: str):
        if not isinstance(color, str) or not color.startswith("#"):
            return None
        hex_value = color.lstrip("#")
        if len(hex_value) == 8:
            hex_value = hex_value[-6:]
        if len(hex_value) != 6:
            return None
        try:
            return tuple(int(hex_value[i:i + 2], 16) for i in (0, 2, 4))
        except ValueError:
            return None

    # Builds a .docx file from our paragraphs, returning the raw file bytes
    def _export_docx(self, paragraphs: list) -> bytes:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from io import BytesIO

        document = Document()
        for runs in paragraphs:
            paragraph = document.add_paragraph()
            for run_data in runs:
                run = paragraph.add_run(run_data.get('text', ""))
                run.bold = bool(run_data.get('bold', False))
                run.italic = bool(run_data.get('italic', False))
                run.underline = bool(run_data.get('underline', False))

                font_size = run_data.get('size')
                if font_size:
                    try:
                        run.font.size = Pt(float(font_size))
                    except (TypeError, ValueError):
                        pass

                font_family = run_data.get('font')
                if font_family:
                    run.font.name = font_family

                rgb = self._hex_to_rgb(run_data.get('color'))
                if rgb:
                    run.font.color.rgb = RGBColor(*rgb)

        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    # Joins our paragraphs into plain text, ready to be written to a .txt file
    def _export_txt(self, paragraphs: list) -> str:
        return "\n".join("".join(run.get('text', "") for run in runs) for runs in paragraphs)

    # Rasterizes our paragraphs onto paginated images and packs them into a multi-page .pdf file
    def _export_pdf(self, paragraphs: list) -> bytes:
        from PIL import Image, ImageDraw, ImageFont
        from io import BytesIO

        font_size = 14
        line_height = int(font_size * 1.5)
        usable_width = MANUSCRIPT_WIDTH - MANUSCRIPT_PADDING * 2
        bottom_limit = MANUSCRIPT_HEIGHT - MANUSCRIPT_PADDING

        font_path = os.path.join(os.path.dirname(__file__), "..", "..", "assets", "fonts", "OpenSans-VariableFont_wdth,wght.ttf")
        try:
            font = ImageFont.truetype(os.path.normpath(font_path), font_size)
        except OSError:
            font = ImageFont.load_default()

        def new_page():
            page = Image.new("RGB", (MANUSCRIPT_WIDTH, MANUSCRIPT_HEIGHT), "white")
            return page, ImageDraw.Draw(page)

        pages = []
        page, draw = new_page()
        cursor_y = MANUSCRIPT_PADDING

        for runs in paragraphs:
            paragraph_text = "".join(run.get('text', "") for run in runs)
            words = paragraph_text.split(" ") if paragraph_text else [""]
            line = ""
            for word in words:
                candidate = f"{line} {word}".strip()
                if line and draw.textlength(candidate, font=font) > usable_width:
                    if cursor_y + line_height > bottom_limit:
                        pages.append(page)
                        page, draw = new_page()
                        cursor_y = MANUSCRIPT_PADDING
                    draw.text((MANUSCRIPT_PADDING, cursor_y), line, font=font, fill="black")
                    cursor_y += line_height
                    line = word
                else:
                    line = candidate

            if cursor_y + line_height > bottom_limit:
                pages.append(page)
                page, draw = new_page()
                cursor_y = MANUSCRIPT_PADDING
            draw.text((MANUSCRIPT_PADDING, cursor_y), line, font=font, fill="black")
            cursor_y += line_height

        pages.append(page)

        buffer = BytesIO()
        pages[0].save(buffer, format="PDF", save_all=True, append_images=pages[1:])
        return buffer.getvalue()
        
        
@ft.component
def ManuscriptView(manuscript: Manuscript) -> WidgetView:

    word_count, set_word_count = ft.use_state(0)

    # Counts words in the editor's current content. Reads the live editor, not the last saved data
    async def refresh_word_count(e=None):
        editor = manuscript.quill_editor
        if editor is None:
            return
        try:
            doc_data = await editor.save()
        except RuntimeError:
            print("Failed to save editor content")
            return  # Editor isn't on the page
        # Join first, so a word split across formatting runs (like a bold letter) counts once. Skips embeds like page breaks
        text = "".join(block["insert"] for block in doc_data if isinstance(block.get("insert"), str))
        set_word_count(len(text.split()))

    # Mark us dirty after keystroke changes so we can track unsaved changes
    def mark_editor_dirty(e=None):
        if manuscript.editor_dirty == False:
            manuscript.editor_dirty = True

    # The popup builds its items when it opens, so they must already be current. Hovering refreshes the count before the click, on_open covers keyboard/touch opens
    word_count_button = ft.GestureDetector(
        ft.PopupMenuButton(
            icon=ft.CupertinoIcons.TEXTFORMAT_SIZE,
            tooltip="Word Count",
            items=[ft.PopupMenuItem(f"Word Count: {word_count:,}", disabled=True, height=20, label_text_style=ft.TextStyle(color=ft.Colors.ON_SURFACE, weight=ft.FontWeight.BOLD))],
            on_open=refresh_word_count,
            style=ft.ButtonStyle(overlay_color=ft.Colors.TRANSPARENT)
        ),
        on_enter=refresh_word_count,
    )
        
    # Grab our flet quill elements
    quill_toolbar = FletQuillToolbar(
        controller_id=manuscript.id,
        expand=True,
        toolbar_buttons=[word_count_button]
    )  
    
    # Flet re-creates controls every render and re-parents only the new instances. A control kept in state keeps a stale parent,
    # so its .page lookup fails and save() raises. Build the editor fresh each render, with a stable key and initial data, and point the manuscript at the newest one
    view_token, _ = ft.use_state(lambda: object())
    initial_data, _ = ft.use_state(lambda: manuscript.manuscript_data)

    quill_editor = FletQuillEditor(
        controller_id=manuscript.id,
        text_data=initial_data,
        placeholder_text="Start your masterpiece here...",
        expand=True,
        key=f"quill_editor_{manuscript.id}"
    )
    manuscript.attach_editor(quill_editor, view_token)

    ft.use_effect(lambda: None, [], cleanup=lambda: manuscript.detach_editor(view_token))

    
    editor_container = ft.Container(
        ft.KeyboardListener(quill_editor, on_key_down=mark_editor_dirty),
        expand=True,
        margin=ft.Margin.symmetric(horizontal=MANUSCRIPT_HORIZONTAL_MARGIN, vertical=MANUSCRIPT_VERTICAL_MARGIN),
    )
    

    return WidgetView(
        manuscript,
        ft.Row([
            ft.Column([
                ft.Container(
                    quill_toolbar, 
                    bgcolor=ft.Colors.SURFACE_CONTAINER_LOWEST, 
                    alignment=ft.Alignment.CENTER_LEFT, 
                    padding=ft.Padding.only(left=4)
                ),
                ft.Divider(2, 2),
                ft.Row([
                    ft.Column([editor_container], scroll=ft.ScrollMode.HIDDEN, expand=True),
                    WidgetSidebar(manuscript), 
                ], spacing=0, expand=True),
            ], spacing=0, expand=True, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
        ], expand=True, spacing=0)
    )